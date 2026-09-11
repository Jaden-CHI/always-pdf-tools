import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor, TimeoutError
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile

from app.adobe_converter import run_adobe_create_pdf, run_adobe_pdf_export
from app.config import Settings


WORD_EXTENSIONS = {".doc", ".docx"}
EXCEL_EXTENSIONS = {".xls", ".xlsx"}
PPT_EXTENSIONS = {".ppt", ".pptx"}
PDF_EXTENSIONS = {".pdf"}


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem.strip() or "converted"
    cleaned = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in stem)
    return cleaned[:80] or "converted"


async def save_upload(file: UploadFile, target: Path, max_bytes: int) -> None:
    total = 0
    with target.open("wb") as output:
        while chunk := await file.read(1024 * 1024):
            total += len(chunk)
            if total > max_bytes:
                raise HTTPException(status_code=413, detail="File is too large.")
            output.write(chunk)


def ensure_extension(filename: str, allowed: set[str]) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in allowed:
        allowed_text = ", ".join(sorted(allowed))
        raise HTTPException(status_code=400, detail=f"Unsupported file type. Allowed: {allowed_text}")
    return suffix


def prepare_office_upload(file: UploadFile, settings: Settings, allowed: set[str]) -> tuple[Path, str]:
    filename = file.filename or "upload"
    suffix = ensure_extension(filename, allowed)
    job_id = uuid4().hex
    job_dir = Path(tempfile.mkdtemp(prefix=f"{job_id}-", dir=settings.work_dir))
    input_path = job_dir / f"input{suffix}"
    return input_path, safe_stem(filename)


def prepare_pdf_upload(file: UploadFile, settings: Settings) -> tuple[Path, str]:
    filename = file.filename or "upload.pdf"
    suffix = ensure_extension(filename, PDF_EXTENSIONS)
    job_id = uuid4().hex
    job_dir = Path(tempfile.mkdtemp(prefix=f"{job_id}-", dir=settings.work_dir))
    input_path = job_dir / f"input{suffix}"
    return input_path, safe_stem(filename)


def run_soffice(input_path: Path, output_stem: str, settings: Settings) -> tuple[Path, str]:
    if settings.adobe_pdf_services_available:
        return run_adobe_create_pdf(input_path, output_stem, settings)

    return run_soffice_to_ext(input_path, output_stem, settings, "pdf")


def run_soffice_to_ext(input_path: Path, output_stem: str, settings: Settings, output_ext: str) -> tuple[Path, str]:
    job_dir = input_path.parent
    command = [
        settings.soffice_path,
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--convert-to",
        output_ext,
        "--outdir",
        str(job_dir),
        str(input_path),
    ]

    try:
        subprocess.run(
            command,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=settings.conversion_timeout_seconds,
        )
    except FileNotFoundError as exc:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(status_code=500, detail="LibreOffice executable was not found.") from exc
    except subprocess.TimeoutExpired as exc:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
    except subprocess.CalledProcessError as exc:
        shutil.rmtree(job_dir, ignore_errors=True)
        stderr = exc.stderr.decode("utf-8", errors="replace")[:500]
        raise HTTPException(status_code=422, detail=f"Conversion failed. {stderr}") from exc

    outputs = sorted(job_dir.glob(f"*.{output_ext}"))
    if not outputs:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(status_code=422, detail=f"Converted {output_ext.upper()} was not created.")

    return outputs[0], f"{output_stem}.{output_ext}"


def _convert_pdf_to_docx(input_path: Path, output_path: Path) -> None:
    from pdf2docx import Converter

    converter = Converter(str(input_path))
    try:
        converter.convert(str(output_path), start=0, end=None)
    finally:
        converter.close()


def run_pdf_to_docx(input_path: Path, output_stem: str, settings: Settings) -> tuple[Path, str]:
    if settings.adobe_pdf_services_available:
        return run_adobe_pdf_export(input_path, output_stem, settings, "docx")

    output_path = input_path.parent / f"{output_stem}.docx"

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(_convert_pdf_to_docx, input_path, output_path)
        try:
            future.result(timeout=settings.conversion_timeout_seconds)
        except TimeoutError as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
        except Exception as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(
                status_code=422,
                detail="PDF to Word conversion failed. Scanned or complex PDFs may not be supported.",
            ) from exc

    if not output_path.exists() or output_path.stat().st_size == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="Converted DOCX was not created.")

    return output_path, f"{output_stem}.docx"


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sheet_title(page_index: int, table_index: int) -> str:
    return f"Page {page_index} Table {table_index}"[:31]


def _convert_pdf_to_xlsx(input_path: Path, output_path: Path) -> int:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    table_count = 0

    with pdfplumber.open(str(input_path)) as pdf:
        for page_index, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table_index, table in enumerate(tables, start=1):
                if not table or not any(any(_normalize_cell(cell) for cell in row) for row in table):
                    continue

                table_count += 1
                sheet = workbook.create_sheet(title=_sheet_title(page_index, table_index))
                for row_index, row in enumerate(table, start=1):
                    for col_index, cell in enumerate(row, start=1):
                        sheet.cell(row=row_index, column=col_index, value=_normalize_cell(cell))

                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="EAF2FF")

                for column_cells in sheet.columns:
                    max_length = max(len(str(cell.value or "")) for cell in column_cells)
                    width = min(max(max_length + 2, 10), 50)
                    sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    if table_count:
        workbook.save(output_path)

    return table_count


def run_pdf_to_xlsx(input_path: Path, output_stem: str, settings: Settings) -> tuple[Path, str]:
    if settings.adobe_pdf_services_available:
        return run_adobe_pdf_export(input_path, output_stem, settings, "xlsx")

    output_path = input_path.parent / f"{output_stem}.xlsx"

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(_convert_pdf_to_xlsx, input_path, output_path)
        try:
            table_count = future.result(timeout=settings.conversion_timeout_seconds)
        except TimeoutError as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
        except Exception as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(
                status_code=422,
                detail="PDF to Excel conversion failed. Complex or scanned PDFs may not be supported.",
            ) from exc

    if table_count == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="No tables were found in this PDF.")

    if not output_path.exists() or output_path.stat().st_size == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="Converted XLSX was not created.")

    return output_path, f"{output_stem}.xlsx"


def _convert_pdf_to_pptx(input_path: Path, output_path: Path) -> None:
    import fitz
    from pptx import Presentation

    document = fitz.open(str(input_path))
    try:
        if document.page_count == 0:
            raise ValueError("PDF has no pages.")

        presentation = Presentation()
        blank_layout = presentation.slide_layouts[6]
        first_page = document.load_page(0)
        presentation.slide_width = int(first_page.rect.width * 12700)
        presentation.slide_height = int(first_page.rect.height * 12700)

        for index in range(document.page_count):
            page = document.load_page(index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_path = input_path.parent / f"page-{index + 1}.png"
            pixmap.save(str(image_path))
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        presentation.save(output_path)
    finally:
        document.close()


def run_pdf_to_pptx(input_path: Path, output_stem: str, settings: Settings) -> tuple[Path, str]:
    if settings.adobe_pdf_services_available:
        return run_adobe_pdf_export(input_path, output_stem, settings, "pptx")

    output_path = input_path.parent / f"{output_stem}.pptx"

    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(_convert_pdf_to_pptx, input_path, output_path)
        try:
            future.result(timeout=settings.conversion_timeout_seconds)
        except TimeoutError as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
        except Exception as exc:
            shutil.rmtree(input_path.parent, ignore_errors=True)
            raise HTTPException(
                status_code=422,
                detail="PDF to PowerPoint conversion failed. Complex or damaged PDFs may not be supported.",
            ) from exc

    if not output_path.exists() or output_path.stat().st_size == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="Converted PPTX was not created.")

    return output_path, f"{output_stem}.pptx"


def run_strong_compress(input_path: Path, output_stem: str, settings: Settings) -> tuple[Path, str]:
    output_path = input_path.parent / f"{output_stem}_compressed.pdf"
    command = [
        settings.gs_path,
        "-sDEVICE=pdfwrite",
        "-dCompatibilityLevel=1.4",
        "-dPDFSETTINGS=/ebook",
        "-dNOPAUSE",
        "-dQUIET",
        "-dBATCH",
        f"-sOutputFile={output_path}",
        str(input_path),
    ]

    try:
        subprocess.run(
            command,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=settings.conversion_timeout_seconds,
        )
    except FileNotFoundError as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=500, detail="Ghostscript executable was not found.") from exc
    except subprocess.TimeoutExpired as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
    except subprocess.CalledProcessError as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        stderr = exc.stderr.decode("utf-8", errors="replace")[:500]
        raise HTTPException(status_code=422, detail=f"Strong compression failed. {stderr}") from exc

    if not output_path.exists() or output_path.stat().st_size == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="Compressed PDF was not created.")

    return output_path, f"{output_stem}_compressed.pdf"


def run_ocrmypdf(input_path: Path, output_stem: str, settings: Settings, mode: str) -> tuple[Path, str]:
    output_path = input_path.parent / f"{output_stem}_{mode}.pdf"
    output_type = "pdfa" if mode == "pdfa" else "pdf"
    command = [
        settings.ocrmypdf_path,
        "--skip-text",
        "--language",
        "kor+eng",
        "--output-type",
        output_type,
        str(input_path),
        str(output_path),
    ]

    try:
        subprocess.run(
            command,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=settings.conversion_timeout_seconds,
        )
    except FileNotFoundError as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=500, detail="OCRmyPDF executable was not found.") from exc
    except subprocess.TimeoutExpired as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=504, detail="Conversion timed out.") from exc
    except subprocess.CalledProcessError as exc:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        stderr = exc.stderr.decode("utf-8", errors="replace")[:500]
        label = "PDF/A" if mode == "pdfa" else "Searchable PDF"
        raise HTTPException(status_code=422, detail=f"{label} conversion failed. {stderr}") from exc

    if not output_path.exists() or output_path.stat().st_size == 0:
        shutil.rmtree(input_path.parent, ignore_errors=True)
        raise HTTPException(status_code=422, detail="Converted PDF was not created.")

    suffix = "pdfa" if mode == "pdfa" else "searchable"
    return output_path, f"{output_stem}_{suffix}.pdf"


def cleanup_file_parent(path: Path) -> None:
    shutil.rmtree(path.parent, ignore_errors=True)
